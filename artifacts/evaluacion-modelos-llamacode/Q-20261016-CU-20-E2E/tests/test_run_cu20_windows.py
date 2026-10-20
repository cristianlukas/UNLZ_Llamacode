"""Regression tests for the isolated CU20 headless runner helpers."""

import pathlib
import copy
import json
import sys
import tempfile
import unittest
from unittest.mock import patch


SCRIPTS = pathlib.Path(__file__).resolve().parents[1] / "scripts"
sys.path.insert(0, str(SCRIPTS))
import run_cu20_windows  # noqa: E402


class TerminalAssistantTests(unittest.TestCase):
    def test_console_json_escapes_text_unsupported_by_windows_code_pages(self):
        line = run_cu20_windows.console_json_line({"title": "Edge\u200b Guest"})
        line.encode("cp1252")
        self.assertEqual(json.loads(line), {"title": "Edge\u200b Guest"})
        self.assertIn(r"\u200b", line)

    def test_final_assistant_is_terminal(self):
        self.assertTrue(run_cu20_windows.has_terminal_assistant([
            {"role": "assistant", "typing": False, "completedAt": 200},
        ]))

    def test_late_appended_earlier_read_does_not_mask_final(self):
        self.assertTrue(run_cu20_windows.has_terminal_assistant([
            {"role": "assistant", "typing": False, "completedAt": 200},
            {"role": "toolcall", "name": "desktop_read_text", "completedAt": 190},
        ]))

    def test_mutation_completed_after_assistant_is_not_terminal(self):
        self.assertFalse(run_cu20_windows.has_terminal_assistant([
            {"role": "assistant", "typing": False, "completedAt": 200},
            {"role": "toolcall", "name": "desktop_type", "completedAt": 210},
        ]))

    def test_user_followup_after_assistant_is_not_terminal(self):
        self.assertFalse(run_cu20_windows.has_terminal_assistant([
            {"role": "assistant", "typing": False, "completedAt": 200},
            {"role": "user", "createdAt": 210},
        ]))

    def test_timestamp_tie_is_not_terminal(self):
        self.assertFalse(run_cu20_windows.has_terminal_assistant([
            {"role": "assistant", "typing": False, "completedAt": 200},
            {"role": "toolcall", "name": "desktop_read_text", "completedAt": 200},
        ]))

    def test_typing_assistant_is_not_terminal(self):
        self.assertFalse(run_cu20_windows.has_terminal_assistant([
            {"role": "assistant", "typing": True, "createdAt": 200},
        ]))

    def test_pending_text_tool_call_is_not_terminal(self):
        self.assertFalse(run_cu20_windows.has_terminal_assistant([
            {"role": "assistant", "typing": False, "completedAt": 200,
             "content": 'TOOL_CALL {"name":"desktop_read_text","arguments":{}}'},
        ]))

    def test_completed_assistant_after_tool_call_is_terminal(self):
        self.assertTrue(run_cu20_windows.has_terminal_assistant([
            {"role": "toolcall", "name": "desktop_read_text", "completedAt": 200},
            {"role": "assistant", "typing": False, "completedAt": 201,
             "content": "La pantalla muestra el aviso exacto."},
        ]))


class CaseFixtureTests(unittest.TestCase):
    def test_cu18_starts_a_clean_session_between_copy_and_paste(self):
        calls = []
        state = {"messages": [{"role": "user"}, {"role": "assistant"}]}

        def fake_invoke(method, args):
            calls.append((method, args))
            state["messages"] = []
            return {"ok": True}

        def fake_prop(name):
            self.assertEqual(name, "agentMessages")
            return state["messages"]

        def fake_wait(predicate, label, timeout_sec):
            self.assertIn("contexto de agente limpio", label)
            self.assertEqual(timeout_sec, 15)
            self.assertTrue(predicate())

        result = run_cu20_windows.start_fresh_agent_session(
            2, fake_invoke, fake_prop, fake_wait)
        self.assertEqual(calls, [("newOpencodeSession", [])])
        self.assertEqual(result, {"method": "newOpencodeSession", "messages_before": 2,
                                  "messages_after": 0, "clean_context": True})
        with self.assertRaisesRegex(ValueError, "conversación previa"):
            run_cu20_windows.start_fresh_agent_session(0, fake_invoke, fake_prop, fake_wait)

    def test_readonly_diagnostic_prompt_only_observes_the_exact_window(self):
        workspace = pathlib.Path("C:/cu20/final-v25-suite-CU13")
        prompt = run_cu20_windows.case_prompt("DIAG_READONLY", workspace,
                                              diagnostic_target_id="518ce")
        self.assertIn("solo lectura", prompt)
        self.assertIn(str(workspace / "cu20-computer-usage.xlsx"), prompt)
        self.assertIn("desktop_read_text dos veces", prompt)
        self.assertIn("query='Sí', query='No', query='Nombre de archivo'", prompt)
        self.assertIn("query='Guardar como'", prompt)
        self.assertIn("No pulses ningún control", prompt)
        self.assertIn("No uses desktop_windows", prompt)
        self.assertIn("DIAG_READONLY", run_cu20_windows.CASE_IDS)

    def test_cu13_save_diagnostic_is_limited_to_the_exact_authorized_dialog(self):
        workspace = pathlib.Path("C:/cu20/final-v25-suite-CU13")
        prompt = run_cu20_windows.case_prompt("DIAG_SAVE_CU13", workspace,
                                              diagnostic_target_id="518ce")
        self.assertIn(str(workspace / "cu20-computer-usage.xlsx"), prompt)
        self.assertIn("autorizó guardar y sobrescribir este libro de prueba", prompt)
        self.assertIn("¿Desea reemplazarlo?", prompt)
        self.assertIn("desktop_controls query='Nombre de archivo'", prompt)
        self.assertIn("desktop_control_action action='read'", prompt)
        self.assertIn("desktop_control_action action='set_value' una sola vez", prompt)
        self.assertIn("value='" + str(workspace / "cu20-computer-usage.xlsx") + "'", prompt)
        self.assertIn("lectura de retorno confirma exactamente ese path", prompt)
        self.assertIn("desktop_controls query='Sí'", prompt)
        self.assertIn("desktop_controls query='Guardar'", prompt)
        self.assertIn("No pulses No, No guardar, Cancelar, Enter ni la X", prompt)
        self.assertIn("sin reenumerar ventanas ni repetir acciones", prompt)
        self.assertIn("DIAG_SAVE_CU13", run_cu20_windows.CASE_IDS)

    def test_test_edge_cleanup_targets_only_the_confirmed_user_profile_window(self):
        prompt = run_cu20_windows.case_prompt(
            "DIAG_CLOSE_TEST_EDGE", pathlib.Path("C:/cu20"),
            diagnostic_target_id="db0f28")
        self.assertIn("target_id exacto db0f28", prompt)
        self.assertIn("PID del Edge normal del usuario", prompt)
        self.assertIn("Microsoft Edge", prompt)
        self.assertIn("Guest/Invitado, InPrivate", prompt)
        self.assertIn("Lee una vez desktop_read_text", prompt)
        self.assertIn("desktop_focus una vez", prompt)
        self.assertIn("desktop_key key='F4', modifiers=['ALT']", prompt)
        self.assertIn("Inmediatamente después ejecuta desktop_windows exactamente una vez", prompt)
        self.assertIn("no hagas ninguna lectura posterior", prompt)
        self.assertIn("Windows puede reciclar ese ID", prompt)
        self.assertIn("si falta db0f28", prompt)
        self.assertIn("no navegues ni cierres ninguna ventana/pestaña previa", prompt)
        self.assertIn("desktop_windows exactamente una vez", prompt)
        self.assertIn("No leas otra fila", prompt)
        self.assertIn("DIAG_CLOSE_TEST_EDGE", run_cu20_windows.CASE_IDS)

    def test_test_excel_discard_targets_only_the_authorized_synthetic_dialog(self):
        workspace = pathlib.Path("C:/cu20/final-v27-suite-CU12")
        prompt = run_cu20_windows.case_prompt(
            "DIAG_DISCARD_TEST_EXCEL", workspace,
            diagnostic_target_id="5450e56", diagnostic_target_pid="24004")
        self.assertIn(str(workspace / "cu20-computer-usage.xlsx"), prompt)
        self.assertIn("PID 24004 y el target_id 5450e56", prompt)
        self.assertIn("¿Guardar los cambios en este archivo?", prompt)
        self.assertIn("desktop_controls query='No guardar'", prompt)
        self.assertIn("desktop_click_element una sola vez", prompt)
        self.assertIn("desktop_windows una vez", prompt)
        self.assertIn("No pulses Guardar, Cancelar, Enter ni la X", prompt)
        self.assertIn("DIAG_DISCARD_TEST_EXCEL", run_cu20_windows.CASE_IDS)
        clipboard_prompt = run_cu20_windows.case_prompt(
            "DIAG_DISCARD_TEST_EXCEL", pathlib.Path("C:/cu20/ae85-cu18fwork"),
            diagnostic_target_id="1a70bb4", diagnostic_target_pid="26736",
            diagnostic_target_file="clipboard-paste.xlsx")
        self.assertIn(str(pathlib.Path("C:/cu20/ae85-cu18fwork") / "clipboard-paste.xlsx"),
                      clipboard_prompt)
        self.assertIn("clipboard-paste.xlsx en Excel", clipboard_prompt)
        self.assertIn("target_id 1a70bb4", clipboard_prompt)

    def test_test_excel_close_uses_only_the_created_pid(self):
        workspace = pathlib.Path("C:/cu20/ae85-cu18cwork")
        prompt = run_cu20_windows.case_prompt(
            "DIAG_CLOSE_TEST_EXCEL", workspace, diagnostic_target_pid="44676",
            diagnostic_target_file="clipboard-paste.xlsx")
        self.assertIn(str(workspace / "clipboard-paste.xlsx"), prompt)
        self.assertIn("PID 44676", prompt)
        self.assertIn("verificado por command line", prompt)
        self.assertIn("clipboard-paste.xlsx - Excel", prompt)
        self.assertIn("desktop_windows una vez", prompt)
        self.assertIn("Enfoca ese target_id", prompt)
        self.assertIn("desktop_controls query='No guardar'", prompt)
        self.assertIn("desktop_key key='F4', modifiers=['ALT']", prompt)
        self.assertNotIn("desktop_key key='S', modifiers=['CTRL']", prompt)
        self.assertIn("no guardes ni pulses Ctrl+S", prompt)
        self.assertIn("otro PID", prompt)

    def test_test_notepad_close_targets_only_the_exact_synthetic_tab(self):
        workspace = pathlib.Path("C:/cu20/final717-cu15work")
        prompt = run_cu20_windows.case_prompt(
            "DIAG_CLOSE_TEST_NOTEPAD", workspace, diagnostic_target_id="31f0b1e")
        self.assertIn("target_id exacto de la ventana Notepad es 31f0b1e", prompt)
        self.assertIn(str(workspace / "notepad-replace.txt"), prompt)
        self.assertIn("desktop_key key='W', modifiers=['CTRL'] una sola vez", prompt)
        self.assertIn("desktop_controls query='No guardar'", prompt)
        self.assertIn("no repitas Ctrl+W", prompt)
        self.assertIn("el título cambió y ya no identifica notepad-replace.txt", prompt)
        self.assertIn("termina ahora y no pulses más teclas", prompt)
        self.assertIn("no cierres otra pestaña", prompt)
        self.assertIn("DIAG_CLOSE_TEST_NOTEPAD", run_cu20_windows.CASE_IDS)
        with tempfile.TemporaryDirectory(prefix="cu14-cleanup-") as temp_dir:
            workspace = pathlib.Path(temp_dir)
            (workspace / "notepad-exact-line.txt").write_text("CU14-EXACT-9F2B", encoding="utf-8")
            prompt = run_cu20_windows.case_prompt(
                "DIAG_CLOSE_TEST_NOTEPAD", workspace, diagnostic_target_id="3810aa")
            self.assertIn("target_id exacto de la ventana Notepad es 3810aa", prompt)
            self.assertIn(str(workspace / "notepad-exact-line.txt"), prompt)
            self.assertIn("título identifica exactamente notepad-exact-line.txt", prompt)
            self.assertIn("no cierres otra pestaña", prompt)
            (workspace / "notepad-no-duplicate.txt").write_text(
                "Synthetic exact-line target.\n", encoding="utf-8")
            cu20_prompt = run_cu20_windows.case_prompt(
                "DIAG_CLOSE_TEST_NOTEPAD", workspace, diagnostic_target_id="2d0a3f",
                diagnostic_target_file="notepad-no-duplicate.txt")
            self.assertIn(str(workspace / "notepad-no-duplicate.txt"), cu20_prompt)
            self.assertIn("título identifica exactamente notepad-no-duplicate.txt", cu20_prompt)

    def test_cu13_prompts_split_read_and_reopen_into_two_turns(self):
        prompt = run_cu20_windows.case_prompt("CU13", pathlib.Path("C:/cu20/workspace"))
        reopen_prompt = run_cu20_windows.cu13_reopen_prompt(pathlib.Path("C:/cu20/workspace"))
        self.assertEqual(prompt.count("Ejecuta desktop_launch una vez"), 1)
        self.assertEqual(reopen_prompt.count("desktop_launch con app="), 1)
        self.assertIn("termina dejando el libro abierto", prompt)
        self.assertIn("El harness ya cerró la ventana anterior", reopen_prompt)
        self.assertIn("Ignora cualquier target_id previo", reopen_prompt)
        self.assertIn("No uses desktop_controls ni hagas más lecturas", prompt)
        self.assertIn("No llames desktop_windows ni desktop_focus", prompt)
        self.assertNotIn("app=`", prompt)
        self.assertNotIn("args=`", prompt)
        self.assertIn("No guardes", prompt)
        self.assertIn("No uses backticks, desktop_windows ni desktop_focus", reopen_prompt)
        self.assertNotIn("app=`", reopen_prompt)
        self.assertNotIn("args=`", reopen_prompt)
        self.assertIn("No edites ni guardes", reopen_prompt)

    def test_cu13_ui_validator_requires_two_turns_and_inter_turn_close(self):
        workbook_args = f'/x "{pathlib.Path("C:/cu20/workspace") / "cu20-computer-usage.xlsx"}"'

        def call(name, args):
            return {"role": "toolcall", "name": name,
                    "arguments": json.dumps(args), "ok": True,
                    "output": ("[desktop_launch: ok — la ventana se abrió]"
                               if name == "desktop_launch" else f"[{name}: ok]")}

        delta = [
            {"role": "user", "content": "stage 1"},
            call("desktop_launch", {"app": str(run_cu20_windows.EXCEL_EXE), "args": workbook_args}),
            call("desktop_read_text", {"query": "cu20-computer-usage.xlsx", "target_id": "w1"}),
            call("desktop_read_text", {"target_id": "w1"}),
            call("desktop_controls", {"target_id": "w1", "query": "Guardar"}),
            call("desktop_read_text", {"target_id": "w1"}),
            {"role": "assistant", "content": "primera lectura lista"},
            {"role": "user", "content": "stage 2"},
            call("desktop_launch", {"app": str(run_cu20_windows.EXCEL_EXE), "args": workbook_args}),
            call("desktop_read_text", {"query": "cu20-computer-usage.xlsx", "target_id": "w2"}),
            call("desktop_read_text", {"target_id": "w2"}),
            call("desktop_controls", {"target_id": "w2", "query": "Guardar"}),
            {"role": "assistant", "content": "segunda lectura lista"},
        ]
        cleanup = {"matched_pids": [1234], "closed_pids": [1234], "remaining_pids": []}
        evidence = run_cu20_windows.validate_ui_evidence("CU13", delta, cleanup)
        self.assertTrue(evidence["passed"])
        self.assertEqual(evidence["launch_count"], 2)
        self.assertEqual(evidence["title_read_count"], 2)
        self.assertEqual(evidence["sheet_read_count"], 3)
        self.assertEqual(evidence["unexpected_action_count"], 0)
        self.assertTrue(evidence["excel_closed_between_turns"])
        self.assertEqual(evidence["user_turn_count"], 2)

        missing_reopen = delta[:5]
        self.assertFalse(run_cu20_windows.validate_ui_evidence("CU13", missing_reopen, cleanup)["passed"])
        already_open = list(delta)
        already_open[8] = {**already_open[8], "output": "[desktop_launch: ok — ya estaba abierta]"}
        self.assertFalse(run_cu20_windows.validate_ui_evidence("CU13", already_open, cleanup)["passed"])
        delayed_reopen = list(delta)
        delayed_reopen.insert(8, call("desktop_windows", {}))
        self.assertFalse(run_cu20_windows.validate_ui_evidence("CU13", delayed_reopen, cleanup)["passed"])
        self.assertFalse(run_cu20_windows.validate_ui_evidence("CU13", delta, None)["passed"])

    def test_cu13_first_turn_requires_launch_and_two_reads(self):
        workbook_args = f'/x "{pathlib.Path("C:/cu20/workspace") / "cu20-computer-usage.xlsx"}"'
        delta = [
            {"role": "user", "content": "stage 1"},
            {"role": "toolcall", "name": "desktop_launch",
             "arguments": json.dumps({"app": str(run_cu20_windows.EXCEL_EXE),
                                      "args": workbook_args}),
             "ok": True, "output": "la ventana se abrió"},
            {"role": "toolcall", "name": "desktop_read_text",
             "arguments": json.dumps({"query": "cu20-computer-usage.xlsx", "target_id": "w1"}),
             "ok": True, "output": "título"},
            {"role": "toolcall", "name": "desktop_read_text",
             "arguments": json.dumps({"target_id": "w1"}), "ok": True, "output": "hoja"},
            {"role": "toolcall", "name": "desktop_controls",
             "arguments": json.dumps({"target_id": "w1", "query": "Guardar"}),
             "ok": True, "output": "controles"},
            {"role": "toolcall", "name": "desktop_read_text",
             "arguments": json.dumps({"target_id": "w1"}), "ok": True, "output": "hoja repetida"},
        ]
        result = run_cu20_windows.validate_cu13_open_turn(delta, pathlib.Path("C:/cu20/workspace"))
        self.assertTrue(result["passed"])
        self.assertEqual(result["launch_count"], 1)
        self.assertEqual(result["sheet_read_count"], 2)
        mutating = list(delta)
        mutating.insert(-1, {"role": "toolcall", "name": "desktop_type",
                             "arguments": json.dumps({"text": "cambio", "target_id": "w1"}),
                             "ok": True, "output": "typed"})
        self.assertFalse(run_cu20_windows.validate_cu13_open_turn(
            mutating, pathlib.Path("C:/cu20/workspace"))["passed"])

    def test_cu13_accepts_only_a_documented_hwnd_replacement(self):
        workbook_args = f'/x "{pathlib.Path("C:/cu20/workspace") / "cu20-computer-usage.xlsx"}"'
        delta = [
            {"role": "user", "content": "stage 1"},
            {"role": "toolcall", "name": "desktop_launch",
             "arguments": json.dumps({"app": str(run_cu20_windows.EXCEL_EXE),
                                      "args": workbook_args}),
             "ok": True, "output": "la ventana se abrió"},
            {"role": "toolcall", "name": "desktop_read_text",
             "arguments": json.dumps({"query": "cu20-computer-usage.xlsx", "target_id": "old"}),
             "ok": True,
             "output": "[desktop_target: el HWND elegido old fue reemplazado por new; seguí usando new]"},
            {"role": "toolcall", "name": "desktop_read_text",
             "arguments": json.dumps({"target_id": "new"}), "ok": True, "output": "hoja"},
        ]
        self.assertTrue(run_cu20_windows.validate_cu13_open_turn(
            delta, pathlib.Path("C:/cu20/workspace"))["passed"])
        delta[2] = {**delta[2], "output": "[desktop_target: HWND not replaced]"}
        self.assertFalse(run_cu20_windows.validate_cu13_open_turn(
            delta, pathlib.Path("C:/cu20/workspace"))["passed"])

    def test_cu11_prompt_uses_editable_name_box_and_persists_one_change(self):
        prompt = run_cu20_windows.case_prompt("CU11", pathlib.Path("C:/cu20/workspace"))
        self.assertIn("no reconstruyas la ruta ni uses backticks ni relances Excel", prompt)
        self.assertIn("botón exacto No una vez", prompt)
        self.assertIn("Ante cualquier otro diálogo, detente sin responderlo", prompt)
        self.assertIn("desktop_controls query='Cuadro de nombres'", prompt)
        self.assertIn("combobox o edit", prompt)
        self.assertIn("control_id, target_id y snapshot_id recientes", prompt)
        self.assertIn("action='set_value', value='B2'", prompt)
        self.assertIn("la siguiente llamada DEBE ser desktop_key key='ENTER', target_id=<Excel>", prompt)
        self.assertIn("No leas ni escribas antes de ese Enter", prompt)
        self.assertIn("desktop_read_text query='B2', target_id=<Excel>, max=8", prompt)
        self.assertIn("la siguiente tool debe ser desktop_read_text query='B2', target_id=<Excel>, max=8", prompt)
        self.assertIn("desktop_type text='4'", prompt)
        self.assertIn("verifica B2=4 y D2=8 con desktop_read_text queries separadas", prompt)
        self.assertIn("No inspecciones ni edites D2", prompt)
        self.assertIn("desktop_controls query='Guardar'", prompt)
        self.assertIn("Exige un único botón Guardar habilitado en la ventana del libro", prompt)
        self.assertIn("desktop_click_element una vez", prompt)
        self.assertIn("No termines antes de guardar", prompt)
        self.assertIn("desktop_read_text query='B2' y desktop_read_text query='D2' por separado", prompt)
        self.assertIn("Si falla, detente", prompt)
        self.assertNotIn("desktop_key key='S', modifiers=['CTRL']", prompt)
        set_value_at = prompt.index("action='set_value', value='B2'")
        enter_at = prompt.index("la siguiente llamada DEBE ser desktop_key key='ENTER'", set_value_at)
        type_at = prompt.index("desktop_type text='4'", enter_at)
        save_at = prompt.index("desktop_controls query='Guardar'", type_at)
        saved_read_at = prompt.index("desktop_read_text query='B2' y desktop_read_text query='D2' por separado", save_at)
        self.assertLess(set_value_at, enter_at)
        self.assertLess(enter_at, type_at)
        self.assertLess(type_at, save_at)
        self.assertLess(save_at, saved_read_at)
        self.assertIn("No uses coordenadas", prompt)
        self.assertIn("detente sin reintentar", prompt)

    def test_every_case_with_semantic_actions_supplies_snapshot_id(self):
        with tempfile.TemporaryDirectory() as directory:
            workspace = pathlib.Path(directory)
            edge_shortcut = workspace / "edge.lnk"
            for case_id in run_cu20_windows.CASE_IDS:
                diagnostic_target_id = (
                    "518ce" if case_id in {"DIAG_READONLY", "DIAG_SAVE_CU13"}
                    else "db0f28" if case_id == "DIAG_CLOSE_TEST_EDGE"
                    else "5450e56" if case_id == "DIAG_DISCARD_TEST_EXCEL"
                    else "31f0b1e" if case_id == "DIAG_CLOSE_TEST_NOTEPAD" else None)
                diagnostic_target_pid = (
                    "24004" if case_id == "DIAG_DISCARD_TEST_EXCEL"
                    else "15928" if case_id == "DIAG_CLOSE_TEST_EXCEL" else None)
                prompt = run_cu20_windows.case_prompt(
                    case_id, workspace, edge_shortcut, diagnostic_target_id, diagnostic_target_pid)
                if "desktop_click_element" in prompt:
                    self.assertIn("snapshot_id", prompt, case_id)
                if "desktop_control_action" in prompt:
                    self.assertIn("snapshot_id", prompt, case_id)

    def test_cu15_prompt_clears_retained_search_state_and_replaces_once(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            prompt = run_cu20_windows.case_prompt("CU15", workspace)
            self.assertEqual(
                (workspace / "notepad-replace.txt").read_text(encoding="utf-8"),
                "The quick brown fox jumps over the sleepy dog.\n")
            self.assertIn("ventana con pestañas", prompt)
            self.assertIn("desktop_key key='H', modifiers=['CTRL']", prompt)
            self.assertIn("desktop_controls query='Buscar'", prompt)
            self.assertIn("desktop_controls query='Reemplazar'", prompt)
            self.assertIn("action='set_value', value='sleepy'", prompt)
            self.assertIn("action='set_value', value='curious'", prompt)
            self.assertIn("desktop_controls query='Reemplazar todo'", prompt)
            self.assertIn("checkpoint de verificación del runtime", prompt)
            self.assertIn("Esa lectura no termina CU15 aunque muestre curious", prompt)
            self.assertIn("No respondas final antes del readback posterior al guardado", prompt)
            self.assertNotIn("desktop_type text=", prompt)

    def test_cu08_prompt_opens_an_empty_tab_with_ctrl_t(self):
        prompt = run_cu20_windows.case_prompt(
            "CU08", pathlib.Path("C:/cu20/workspace"), pathlib.Path("C:/cu20/edge.lnk"))
        self.assertIn("desktop_key key='T', modifiers=['CTRL']", prompt)
        self.assertIn("Apollo 12 landing date", prompt)
        self.assertIn("confirma la pestaña nueva", prompt)
        self.assertIn("No cierres pestañas ni repitas búsquedas", prompt)
        self.assertIn("target_id=<Edge>", prompt)
        self.assertIn("desktop_key key='ENTER', target_id=<Edge>", prompt)

    def test_cu07_prompt_submits_navigation_and_types_find_query(self):
        prompt = run_cu20_windows.case_prompt(
            "CU07", pathlib.Path("C:/cu20/workspace"), pathlib.Path("C:/cu20/edge.lnk"))
        for required in ("desktop_launch con app igual al literal exacto",
                         "args como cadena vacía",
                         "desktop_key key='L', modifiers=['CTRL'], target_id=<Edge>",
                         "lee query='Apollo 11', max=4",
                         "desktop_key key='F', modifiers=['CTRL'], target_id=<Edge>",
                         "query='Buscar en la página', max=3",
                         "desktop_type text='July 20, 1969', target_id=<Edge>",
                         "query='July 20, 1969', max=8",
                         "Confirma la fecha en OCR del artículo"):
            self.assertIn(required, prompt)
        self.assertLess(prompt.index("desktop_key key='ENTER'"),
                        prompt.index("desktop_key key='F'"))

    def test_edge_shortcut_skips_first_run_and_reuses_current_user_profile(self):
        arguments = run_cu20_windows.edge_shortcut_arguments()
        self.assertNotIn("--user-data-dir", arguments)
        self.assertNotIn("--guest", arguments.casefold())
        self.assertNotIn("--inprivate", arguments.casefold())
        self.assertIn("--new-window", arguments)
        self.assertIn("--no-first-run", arguments)

    def test_edge_launch_requires_the_user_profile_shortcut_and_a_new_window(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            shortcut = pathlib.Path(temp_dir) / "Edge-CU20-UserProfile.lnk"
            messages = [{
                "role": "toolcall", "name": "desktop_launch", "ok": True,
                "arguments": json.dumps({"app": str(shortcut), "args": ""}),
                "output": "desktop_launch: ok receipt_id=abc [desktop_launch: ok — "
                          "la ventana se abrió: id=1a2b3c4 \"Nueva pestaña\"]",
            }]
            run_cu20_windows.enforce_edge_user_profile_launch(messages, shortcut)

    def test_edge_launch_rejects_reusing_an_existing_user_window(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            shortcut = pathlib.Path(temp_dir) / "Edge-CU20-UserProfile.lnk"
            messages = [{
                "role": "toolcall", "name": "desktop_launch", "ok": True,
                "arguments": json.dumps({"app": str(shortcut), "args": ""}),
                "output": "desktop_launch: ok receipt_id=abc [desktop_launch: ok — "
                          "la ventana ya estaba abierta (no apareció otra): id=1a2b3c4 "
                          "\"New tab - Microsoft Edge\"]",
            }]
            with self.assertRaises(run_cu20_windows.TurnSafetyError):
                run_cu20_windows.enforce_edge_user_profile_launch(messages, shortcut)

    def test_excel_recovery_authorization_is_limited_to_synthetic_book_cases(self):
        prompt = "Task prompt"
        workspace = pathlib.Path(r"C:\cu20\synthetic-workspace")
        excel = run_cu20_windows.apply_excel_recovery_authorization("CU18", prompt, workspace)
        self.assertIn("autorizó cerrar aplicaciones y descartar cambios", excel)
        self.assertIn(str(workspace), excel)
        self.assertIn("cu20-computer-usage.xlsx, clipboard-paste.xlsx", excel)
        self.assertIn("No guardar", excel)
        self.assertIn("No descartes documentos ajenos", excel)
        self.assertEqual(
            run_cu20_windows.apply_excel_recovery_authorization("CU05", prompt), prompt)

    def test_cu16_prompt_observes_and_uses_window_relative_canvas_points(self):
        prompt = run_cu20_windows.case_prompt(
            "CU16", pathlib.Path("C:/cu20/workspace"), None)
        self.assertIn("desktop_observe", prompt)
        self.assertEqual(prompt.count("desktop_observe"), 1)
        self.assertNotIn("desktop_snapshot", prompt)
        self.assertIn("lienzo en blanco", prompt)
        self.assertIn("x=0.279, y=0.088", prompt)
        self.assertIn('{"x":0.44,"y":0.47}', prompt)
        self.assertIn('{"x":0.56,"y":0.66}', prompt)
        self.assertIn('[{"x":0.44,"y":0.47},{"x":0.56,"y":0.66}]', prompt)
        self.assertIn("NO uses pares internos [[0.44,0.47],[0.56,0.66]]", prompt)
        self.assertIn("el harness comprobará independientemente", prompt)
        self.assertIn("No uses UIA ni repitas el clic", prompt)
        self.assertNotIn("desktop_observe otra vez", prompt)
        self.assertNotIn("desktop_click_element", prompt)

    def test_cu09_prepares_empty_workbook_and_hash(self):
        from openpyxl import load_workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            fixture = run_cu20_windows.prepare_case_fixture("CU09", workspace)
            self.assertEqual(fixture["kind"], "empty_xlsx")
            self.assertEqual(len(fixture["sha256_before_agent"]), 64)
            workbook = load_workbook(fixture["path"], data_only=False)
            self.assertIsNone(workbook.active["A1"].value)
            workbook.close()

    def test_cu09_prompt_quotes_and_repeats_the_exact_workbook_argument(self):
        workspace = pathlib.Path("C:/work space/CU20")
        expected = '/x "' + str(workspace / "cu20-computer-usage.xlsx") + '"'
        for case_id in ("CU09", "CU10", "CU11", "CU12", "CU13"):
            with self.subTest(case_id=case_id):
                prompt = run_cu20_windows.case_prompt(case_id, workspace)
                if case_id == "CU10":
                    self.assertIn(f"app={str(run_cu20_windows.EXCEL_EXE)!r}", prompt)
                    self.assertIn(f"args={expected!r}", prompt)
                    self.assertIn("No uses backticks en los valores", prompt)
                    self.assertNotIn("app=`", prompt)
                    self.assertNotIn("args=`", prompt)
                elif case_id in {"CU11", "CU13"}:
                    self.assertIn(f"app={str(run_cu20_windows.EXCEL_EXE)!r}", prompt)
                    self.assertIn(f"args={expected!r}", prompt)
                    self.assertIn("backticks", prompt)
                    self.assertNotIn("app=`", prompt)
                    self.assertNotIn("args=`", prompt)
                else:
                    self.assertIn(repr(expected), prompt)
                self.assertIn("reconstruyas", prompt)
        prompt = run_cu20_windows.case_prompt("CU09", workspace)
        self.assertIn("confirma por el título que abriste cu20-computer-usage.xlsx", prompt)

    def test_cu12_prompt_requires_both_formats_and_save_evidence(self):
        workspace = pathlib.Path("C:/work space/CU20")
        prompt = run_cu20_windows.case_prompt("CU12", workspace)
        self.assertIn("desktop_read_text query='cu20-computer-usage.xlsx', target_id=<Excel>, max=8", prompt)
        self.assertIn("desktop_wait_for window_title='cu20-computer-usage.xlsx', timeout_ms=30000", prompt)
        self.assertIn("No llames desktop_windows, desktop_focus ni desktop_controls antes", prompt)
        self.assertIn("Cuadro de nombres", prompt)
        self.assertIn("action='set_value', value='A1:D1'", prompt)
        self.assertIn("desktop_controls query='Negrita'", prompt)
        self.assertIn("desktop_click_element una vez", prompt)
        self.assertIn("action='set_value', value='C2:C4'", prompt)
        self.assertIn("aplica formato moneda una sola vez con desktop_key key='4'", prompt)
        self.assertIn("desktop_controls query='Guardar'", prompt)
        self.assertIn("No uses Ctrl+B ni Ctrl+S", prompt)
        self.assertIn("detente sin repetirla", prompt)
        self.assertIn("Inmediatamente lee desktop_read_text query='Cuadro de nombres'", prompt)
        self.assertIn("Cuadro de nombres: A1:D1", prompt)
        self.assertIn("action='set_value', value='C2:C4' una vez", prompt)
        self.assertIn("Cuadro de nombres: C2:C4", prompt)
        self.assertIn("no interpretes 0 renglones como fallo", prompt)
        self.assertIn("no vuelvas a seleccionar el rango", prompt)
        self.assertNotIn("desktop_read_text query='C2:C4'", prompt)
        self.assertIn("aplica formato moneda una sola vez con desktop_key key='4'", prompt)
        self.assertNotIn("ni sustituirla por coordenadas, texto o atajos", prompt)

    def test_cu05_prompt_requires_subject_in_evidence_based_answer(self):
        prompt = run_cu20_windows.case_prompt(
            "CU05", pathlib.Path("C:/work space/CU20"),
            pathlib.Path("C:/temp/Edge-CU20-UserProfile.lnk"))
        self.assertIn("menciona explícitamente Apollo 11 junto con la fecha", prompt)
        self.assertIn("No declares éxito si el resultado visible no respalda", prompt)
        self.assertIn("desktop_wait_for window_title='Apollo 11' timeout_ms=20000 una vez", prompt)
        self.assertIn("inmediatamente después ejecuta una desktop_read_text nueva", prompt)
        self.assertIn("query='Apollo 11 landing date', max=8", prompt)
        self.assertIn("Enter una sola vez únicamente si esa lectura confirma el texto", prompt)
        self.assertIn("perfil existente del usuario", prompt)
        self.assertIn("no abras Guest/Invitado ni InPrivate", prompt)
        self.assertIn("Nunca navegues ni cierres pestañas/ventanas que ya estaban abiertas", prompt)
        self.assertIn("aparece un CAPTCHA", prompt)
        self.assertIn("Después de Enter ejecuta inmediatamente desktop_read_text", prompt)
        self.assertLess(prompt.index("key='L', modifiers=['CTRL']"),
                        prompt.index("desktop_read_text nueva"))
        self.assertLess(prompt.index("Apollo 11 landing date una sola vez"),
                        prompt.index("query='Apollo 11 landing date', max=8"))
        self.assertLess(prompt.index("query='Apollo 11 landing date', max=8"),
                        prompt.index("Enter una sola vez únicamente"))

    def test_cu06_prompt_observes_each_browser_mutation(self):
        prompt = run_cu20_windows.case_prompt(
            "CU06", pathlib.Path("C:/work space/CU20"),
            pathlib.Path("C:/temp/Edge-CU20-UserProfile.lnk"))
        self.assertIn("target_id exacto de desktop_launch", prompt)
        self.assertIn("desktop_read_text target_id=<Edge>, max=8", prompt)
        self.assertIn("https://www.google.com/search?q=historia+de+Egipto+en+Wikipedia", prompt)
        self.assertIn("query='historia+de+Egipto+en+Wikipedia', max=8", prompt)
        self.assertIn("Pulsa Enter una vez sólo si confirma la URL", prompt)
        self.assertIn("query='Wikipedia', max=12", prompt)
        self.assertIn("Historia de Egipto'", prompt)
        self.assertIn("snippet del resultado de Wikipedia", prompt)
        self.assertIn("encabezado del artículo", prompt)
        self.assertIn("centro normalizado (x,y) de esa misma línea", prompt)
        self.assertIn("desktop_click scope_kind='window'", prompt)
        self.assertIn("Si Google muestra CAPTCHA", prompt)
        self.assertIn("no esperes ni repitas lecturas/búsquedas", prompt)
        self.assertIn("No inventes coordenadas", prompt)
        self.assertNotIn("desktop_click_element", prompt)
        self.assertLess(prompt.index("modifiers=['CTRL']"),
                        prompt.index("desktop_read_text target_id=<Edge>, max=8"))

    def test_cu07_prompt_observes_url_and_find_query_before_enter(self):
        shortcut = pathlib.Path("C:/temp/Edge-CU20-UserProfile.lnk")
        prompt = run_cu20_windows.case_prompt(
            "CU07", pathlib.Path("C:/work space/CU20"),
            shortcut)
        self.assertIn("target_id exacto", prompt)
        self.assertIn("literal exacto", prompt)
        self.assertIn(repr(str(shortcut)), prompt)
        self.assertIn("args como cadena vacía", prompt)
        self.assertIn("no reconstruyas ni cambies letras o números", prompt)
        self.assertIn("query='wikipedia', max=3", prompt)
        self.assertIn("desktop_key key='ENTER', target_id=<Edge> una vez", prompt)
        self.assertIn("query='Buscar en la página', max=3", prompt)
        self.assertIn("query='July 20, 1969', max=8", prompt)

    def test_cu08_prompt_observes_each_search_and_tab_change(self):
        prompt = run_cu20_windows.case_prompt(
            "CU08", pathlib.Path("C:/work space/CU20"),
            pathlib.Path("C:/temp/Edge-CU20-UserProfile.lnk"))
        self.assertIn("query='Apollo 11 landing date', max=3", prompt)
        self.assertIn("query='Apollo 12 landing date', max=3", prompt)
        self.assertIn("query='Apollo 11', max=3 para confirmar el regreso", prompt)
        self.assertIn("Conserva el target_id exacto que devuelve y úsalo en", prompt)
        self.assertIn("perfil normal existente del usuario", prompt)

    def test_cu07_prompt_observes_find_panel_before_typing(self):
        prompt = run_cu20_windows.case_prompt(
            "CU07", pathlib.Path("C:/work space/CU20"),
            pathlib.Path("C:/temp/Edge-CU20-UserProfile.lnk"))
        self.assertIn("lee query='Apollo 11', max=4", prompt)
        self.assertIn("query='Buscar en la página', max=3", prompt)
        self.assertIn("query='July 20, 1969', max=4", prompt)
        self.assertIn("Confirma la fecha en OCR del artículo", prompt)
        self.assertIn("Si esa lectura ya muestra la coincidencia en el OCR, no pulses Enter", prompt)
        self.assertLess(prompt.index("desktop_key key='F'"),
                        prompt.index("desktop_type text='July 20, 1969'"))

    def test_cu18_prompt_quotes_and_repeats_the_exact_workbook_argument(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir) / "work space" / "CU20"
            workspace.mkdir(parents=True)
            expected = '/x "' + str(workspace / "clipboard-paste.xlsx") + '"'
            prompt = run_cu20_windows.cu18_paste_prompt(workspace)
            self.assertIn("args=" + repr(expected), prompt)
            self.assertIn("no uses backticks", prompt)
            self.assertNotIn("args=`", prompt)
            self.assertIn("conserva /x y las comillas de la ruta", prompt)
            self.assertIn("desktop_controls query='Cuadro de nombres'", prompt)
            self.assertIn("desktop_key key='V', modifiers=['CTRL']", prompt)
            self.assertIn(f"app={str(run_cu20_windows.EXCEL_EXE)!r}", prompt)

    def test_excel_process_matcher_requires_exact_test_workbook_path(self):
        workspace = pathlib.Path("C:/cu20/workspace")
        workbook = workspace / "cu20-computer-usage.xlsx"
        matching = '"C:/Program Files/Microsoft Office/EXCEL.EXE" /x "C:/cu20/workspace/cu20-computer-usage.xlsx"'
        unrelated = '"C:/Program Files/Microsoft Office/EXCEL.EXE" /x C:/Users/user/Book1.xlsx'
        similar = '"C:/Program Files/Microsoft Office/EXCEL.EXE" /x C:/cu20/workspace/cu20-computer-usage.xlsx.bak'
        self.assertTrue(run_cu20_windows.is_test_excel_process("EXCEL.EXE", matching, workbook))
        self.assertFalse(run_cu20_windows.is_test_excel_process("EXCEL.EXE", unrelated, workbook))
        self.assertFalse(run_cu20_windows.is_test_excel_process("EXCEL.EXE", similar, workbook))
        self.assertFalse(run_cu20_windows.is_test_excel_process("notepad.exe", matching, workbook))

    def test_excel_cleanup_requests_graceful_close_without_kill(self):
        class FakeProcess:
            pid = 321
            info = {
                "name": "EXCEL.EXE",
                "cmdline": ["EXCEL.EXE", "/x", "C:/cu20/workspace/cu20-computer-usage.xlsx"],
            }

            def is_running(self):
                return True

            def terminate(self):
                raise AssertionError("cleanup must not terminate Excel")

            def kill(self):
                raise AssertionError("cleanup must not kill Excel")

        process = FakeProcess()
        workspace = pathlib.Path("C:/cu20/workspace")
        with patch("psutil.process_iter", return_value=[process]), \
                patch("psutil.wait_procs", return_value=([], [process])), \
                patch.object(run_cu20_windows, "request_close_test_excel_windows",
                             return_value=[1234]) as request_close, \
                patch.object(run_cu20_windows, "visible_test_excel_window_pids",
                             return_value={321}):
            result = run_cu20_windows.close_test_excel_processes("CU10", workspace)

        request_close.assert_called_once_with({321}, "cu20-computer-usage.xlsx")
        self.assertEqual(result["close_requested_windows"], [1234])
        self.assertEqual(result["remaining_pids"], [321])

    def test_excel_cleanup_does_not_block_on_process_without_workbook_window(self):
        class FakeProcess:
            pid = 321
            info = {
                "name": "EXCEL.EXE",
                "cmdline": ["EXCEL.EXE", "/x", "C:/cu20/workspace/cu20-computer-usage.xlsx"],
            }

            def is_running(self):
                return True

        process = FakeProcess()
        workspace = pathlib.Path("C:/cu20/workspace")
        with patch("psutil.process_iter", return_value=[process]), \
                patch("psutil.wait_procs", return_value=([], [process])), \
                patch.object(run_cu20_windows, "request_close_test_excel_windows",
                             return_value=[1234]), \
                patch.object(run_cu20_windows, "visible_test_excel_window_pids",
                             return_value=set()):
            result = run_cu20_windows.close_test_excel_processes("CU10", workspace)

        self.assertEqual(result["remaining_pids"], [])
        self.assertEqual(result["processes_alive_after_close"], [321])

    def test_excel_cleanup_with_no_match_is_noop(self):
        with patch("psutil.process_iter", return_value=[]), \
                patch.object(run_cu20_windows, "request_close_test_excel_windows") as request_close:
            result = run_cu20_windows.close_test_excel_processes(
                "CU10", pathlib.Path("C:/cu20/workspace"))

        request_close.assert_called_once_with(set(), "cu20-computer-usage.xlsx")
        self.assertEqual(result["matched_pids"], [])
        self.assertEqual(result["remaining_pids"], [])

    def test_cu09_refuses_populated_workspace(self):
        from openpyxl import Workbook

        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            workbook = Workbook()
            workbook.active["A1"] = "stale test data"
            workbook.save(workspace / "cu20-computer-usage.xlsx")
            workbook.close()
            with self.assertRaisesRegex(RuntimeError, "workspace nuevo"):
                run_cu20_windows.prepare_case_fixture("CU09", workspace)

    def test_excel_cases_create_their_expected_seeded_fixture(self):
        from openpyxl import load_workbook

        expected = {
            "CU10": (2, None, None, False, False),
            "CU11": (2, "Total cantidades", "=B2+B3+B4", False, False),
            "CU12": (4, "Total cantidades", "=B2+B3+B4", False, False),
            "CU13": (4, "Total cantidades", "=B2+B3+B4", True, True),
        }
        for case_id, state in expected.items():
            with self.subTest(case_id=case_id), tempfile.TemporaryDirectory() as temp_dir:
                workspace = pathlib.Path(temp_dir)
                fixture = run_cu20_windows.prepare_case_fixture(case_id, workspace)
                self.assertEqual(fixture["kind"], "seeded_xlsx")
                self.assertEqual(fixture["case_id"], case_id)
                self.assertEqual(len(fixture["sha256_before_agent"]), 64)
                workbook = load_workbook(fixture["path"], data_only=False)
                sheet = workbook.active
                self.assertEqual(sheet["B2"].value, state[0])
                self.assertEqual(sheet["D1"].value, state[1])
                self.assertEqual(sheet["D2"].value, state[2])
                self.assertEqual(all(sheet.cell(1, col).font.bold for col in range(1, 5)), state[3])
                self.assertEqual(all("$" in sheet.cell(row, 3).number_format
                                     for row in range(2, 5)), state[4])
                self.assertEqual([[sheet.cell(row, col).value for col in range(1, 4)]
                                  for row in range(1, 5)],
                                 [["Producto", "Cantidad", "Precio"],
                                  ["Café", state[0], 1250], ["Té", 3, 800],
                                  ["Agua", 1, 500]])
                workbook.close()
                if case_id == "CU13":
                    cached_workbook = load_workbook(fixture["path"], data_only=True)
                    self.assertEqual(cached_workbook.active["D2"].value, 8)
                    cached_workbook.close()

    def test_other_cases_do_not_create_excel_fixture(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            self.assertEqual(run_cu20_windows.prepare_case_fixture("CU15", workspace), {})
            self.assertFalse((workspace / "cu20-computer-usage.xlsx").exists())

    def test_cu18_prepares_blank_clipboard_destination(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            fixture = run_cu20_windows.prepare_case_fixture("CU18", workspace)
            self.assertEqual(fixture["case_id"], "CU18")
            self.assertEqual(pathlib.Path(fixture["path"]).name, "clipboard-paste.xlsx")
            self.assertEqual(len(fixture["sha256_before_agent"]), 64)

    def test_cu16_prepares_blank_png_in_new_workspace(self):
        from PIL import Image

        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            fixture = run_cu20_windows.prepare_case_fixture("CU16", workspace)
            self.assertEqual(fixture["kind"], "blank_png")
            self.assertEqual(fixture["size"], [400, 300])
            self.assertEqual(len(fixture["sha256_before_agent"]), 64)
            with Image.open(fixture["path"]) as image:
                self.assertEqual(image.format, "PNG")
                self.assertEqual(image.size, (400, 300))
                self.assertEqual(image.convert("RGB").getpixel((200, 150)), (255, 255, 255))

    def test_cu16_refuses_existing_workspace_png(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            (workspace / "paint-rectangle.png").write_bytes(b"preserve this file")
            with self.assertRaisesRegex(RuntimeError, "workspace nuevo"):
                run_cu20_windows.prepare_case_fixture("CU16", workspace)

    def test_cu18_prompt_uses_single_semantic_copy_paste_and_exact_source(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            prompt = run_cu20_windows.case_prompt("CU18", workspace)
            source = workspace / "clipboard-source.txt"
            self.assertEqual(source.read_text(encoding="utf-8"),
                             run_cu20_windows.CLIPBOARD_MARKER)
            self.assertIn("query='CU20-CLIPBOARD-7K9-PRESERVE'", prompt)
            self.assertIn("desktop_focus", prompt)
            self.assertIn("desktop_read_text query=", prompt)
            self.assertIn("Después enfoca ese mismo target_id", prompt)
            self.assertIn("desktop_key no recibe target_id", prompt)
            self.assertIn("desktop_key key='A', modifiers=['CTRL']", prompt)
            self.assertIn("DEBE ser desktop_key key='C', modifiers=['CTRL']", prompt)
            self.assertIn("no abras Excel", prompt)
            self.assertIn("desktop_clipboard op='get'", prompt)
            self.assertNotIn(str(run_cu20_windows.EXCEL_EXE), prompt)
            paste_prompt = run_cu20_windows.cu18_paste_prompt(workspace)
            self.assertIn("Abre una vez Excel con app=", paste_prompt)
            self.assertIn("desktop_wait_for window_title='clipboard-paste.xlsx - Excel'", paste_prompt)
            self.assertIn("desktop_read_text query='clipboard-paste.xlsx'", paste_prompt)
            self.assertIn("desktop_controls query='Cuadro de nombres'", paste_prompt)
            self.assertIn("desktop_controls query='Guardar'", paste_prompt)
            self.assertIn("desktop_control_action action='set_value', value='A1'", paste_prompt)
            self.assertIn("control_id, target_id y snapshot_id recientes", paste_prompt)
            self.assertIn("desktop_click_element", paste_prompt)
            self.assertNotIn("desktop_key key='S'", paste_prompt)
            self.assertIn("desktop_read_text con sólo target_id", paste_prompt)
            self.assertIn("Guardado en Este PC", paste_prompt)
            self.assertIn(f"app={str(run_cu20_windows.EXCEL_EXE)!r}", paste_prompt)
            self.assertNotIn("app=`", paste_prompt)
            self.assertIn("desktop_key key='V', modifiers=['CTRL']", paste_prompt)
            self.assertIn("No uses Ctrl+S", paste_prompt)
            self.assertIn("detente sin repetir el pegado", paste_prompt)
            expected_excel_args = '/x "' + str(workspace / 'clipboard-paste.xlsx') + '"'
            self.assertIn("args=" + repr(expected_excel_args), paste_prompt)
            self.assertNotIn("args=`", paste_prompt)
            self.assertNotIn("desktop_windows", paste_prompt)
            self.assertNotIn("apply_excel_recovery_authorization", paste_prompt)

    def test_cu14_prompt_targets_only_the_synthetic_empty_file(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            prompt = run_cu20_windows.case_prompt("CU14", workspace)
            path = workspace / "notepad-exact-line.txt"
            self.assertTrue(path.is_file())
            self.assertEqual(path.read_text(encoding="utf-8"), "")
            self.assertIn("archivo sintético", prompt)
            self.assertIn("desktop_wait_for window_title='notepad-exact-line.txt'", prompt)
            self.assertIn("desktop_read_text query='notepad-exact-line.txt', max=8", prompt)
            self.assertIn("Continúa sólo si la ventana confirma ese nombre exacto", prompt)
            self.assertIn("Escribe una sola vez", prompt)
            self.assertIn("desktop_key key='S', modifiers=['CTRL']", prompt)
            self.assertIn("desktop_read_text query=", prompt)
            self.assertIn("No abras otra app", prompt)

    def test_cu16_requires_four_axis_aligned_sides(self):
        from PIL import Image, ImageDraw

        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            path = workspace / "paint-rectangle.png"
            image = Image.new("RGB", (400, 300), "white")
            ImageDraw.Draw(image).rectangle((70, 50, 330, 250), outline="black", width=3)
            image.save(path)
            rectangle = run_cu20_windows.validate_artifact("CU16", workspace)
            self.assertTrue(rectangle["passed"], rectangle)
            self.assertTrue(rectangle["rectangular_outline"])

            image = Image.new("RGB", (400, 300), "white")
            ImageDraw.Draw(image).line((70, 50, 330, 250), fill="black", width=3)
            image.save(path)
            diagonal = run_cu20_windows.validate_artifact("CU16", workspace)
            self.assertFalse(diagonal["passed"], diagonal)
            self.assertFalse(diagonal["rectangular_outline"])


class UiEvidenceTests(unittest.TestCase):
    @staticmethod
    def tool(name, output, arguments="{}"):
        return {"role": "toolcall", "name": name, "ok": True,
                "arguments": arguments, "output": output}

    def test_cu19_prompt_opens_two_unique_notepad_documents(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir) / "work space"
            workspace.mkdir()
            alpha, beta = run_cu20_windows.cu19_document_paths(workspace)
            prompt = run_cu20_windows.case_prompt("CU19", workspace)
            self.assertIn("desktop_launch app='notepad'", prompt)
            self.assertIn('args=\'"' + str(alpha) + '"\'', prompt)
            self.assertIn("desktop_focus", prompt)
            self.assertIn("desktop_windows", prompt)
            self.assertIn("Inmediatamente lee una vez desktop_read_text sin query", prompt)
            self.assertIn("un cuerpo vacío es normal y no bloquea la prueba", prompt)
            self.assertIn("query='Editor de texto', max=4", prompt)
            self.assertIn("keyboardFocusable=1", prompt)
            self.assertIn("keyboardFocusConfirmed=true", prompt)
            self.assertEqual(prompt.count("desktop_launch app='notepad'"), 2)
            self.assertIn(str(beta), prompt)
            self.assertIn(alpha.name, prompt)
            self.assertIn(beta.name, prompt)
            self.assertIn("No uses el texto del documento como prueba del título", prompt)
            self.assertGreaterEqual(prompt.count("query='contenido beta solicitado'"), 2)
            self.assertIn("si no coincide exactamente, detente sin guardar", prompt)
            self.assertIn("desktop_key key='S', modifiers=['CTRL']", prompt)
            self.assertNotIn("desktop_launch app='notepad.exe'", prompt)

    def test_cu20_prompt_confirms_and_targets_existing_notepad_before_mutation(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir) / "final-CU20"
            workspace.mkdir()
            prompt = run_cu20_windows.case_prompt("CU20", workspace)
            self.assertIn("desktop_read_text target_id=<id devuelto>", prompt)
            self.assertIn("desktop_focus", prompt)
            self.assertIn("desktop_type text=<frase>, target_id=<id confirmado>", prompt)
            self.assertIn("desktop_key key='S', modifiers=['CTRL']", prompt)
            self.assertIn("key lleva 'S', nunca 'CTRL'", prompt)
            self.assertIn("desktop_read_text y confirma que la frase sigue apareciendo una sola vez", prompt)

    def test_cu19_exposes_only_tools_needed_for_the_sequence(self):
        catalog = [
            "desktop_assert", "desktop_click_element", "desktop_clipboard",
            "desktop_control_action", "desktop_controls", "desktop_focus",
            "desktop_key", "desktop_launch", "desktop_read_text", "desktop_type",
            "desktop_windows",
        ]
        selected = run_cu20_windows.tools_for_case("CU19", catalog)
        self.assertEqual(selected, sorted(run_cu20_windows.CASE_TOOL_ALLOWLIST["CU19"]))
        self.assertNotIn("desktop_clipboard", selected)
        self.assertEqual(run_cu20_windows.tools_for_case("CU14", catalog), sorted(catalog))
        with self.assertRaises(RuntimeError):
            run_cu20_windows.tools_for_case("CU19", ["desktop_launch"])

    def test_cu19_requires_two_title_confirmations_and_only_writes_beta(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            alpha, beta = run_cu20_windows.cu19_document_paths(workspace)
            # Notepad may switch documents as tabs under one top-level HWND.
            alpha_id = beta_id = "a1b2c3"

            def tool(name, arguments, output="ok", ok=True):
                return {"role": "toolcall", "name": name, "ok": ok,
                        "arguments": json.dumps(arguments), "output": output}

            delta = [
                tool("desktop_launch", {"app": "notepad", "args": f'"{alpha}"'},
                     f"[desktop_launch: ok id={alpha_id} opened]"),
                tool("desktop_read_text", {"target_id": alpha_id},
                     f"[desktop_read_text: window '{alpha_id}']\n"
                     f"Título de ventana: {alpha.name}: Bloc de notas"),
                tool("desktop_launch", {"app": "notepad", "args": f'"{beta}"'},
                     f"[desktop_launch: ok id={beta_id} opened]"),
                tool("desktop_read_text", {"target_id": beta_id},
                     f"[desktop_read_text: window '{beta_id}']\n"
                     f"Título de ventana: {beta.name}: Bloc de notas\n"
                     f"Texto de la app (UI Automation):\n{alpha.name}\n{beta.name}"),
                tool("desktop_windows", {},
                     f'id={beta_id} pid=123 900x600 "{beta.name}: Bloc de notas"'),
                tool("desktop_focus", {"target_id": beta_id},
                     "[desktop_focus: ok]"),
                tool("desktop_controls", {"target_id": beta_id,
                                          "query": "Editor de texto", "max": 4},
                     'snapshot_id=snap-editor-1\n'
                     'controlId=editor-doc [document] focusable=1 keyboard_focus=1 "Editor de texto"'),
                tool("desktop_click_element", {"target_id": beta_id,
                                                  "control_id": "editor-doc",
                                                  "snapshot_id": "snap-editor-1"},
                     'trace=' + json.dumps({"action": "click_element",
                                            "keyboardFocusConfirmed": True,
                                            "target": {"role": "document"}})),
                tool("desktop_type", {"target_id": beta_id, "text": "contenido beta solicitado"}),
                tool("desktop_read_text", {"target_id": beta_id,
                                            "query": "contenido beta solicitado"},
                     f"Título de ventana: {beta.name}: Bloc de notas\ncontenido beta solicitado"),
                tool("desktop_key", {"key": "S", "modifiers": ["CTRL"]}),
                tool("desktop_read_text", {"target_id": beta_id,
                                            "query": "contenido beta solicitado"},
                     f"Título de ventana: {beta.name}: Bloc de notas\ncontenido beta solicitado"),
            ]
            evidence = run_cu20_windows.validate_ui_evidence(
                "CU19", delta, expected_workspace=workspace)
            self.assertTrue(evidence["passed"], evidence)

            broad_controls = copy.deepcopy(delta)
            broad_controls[6]["arguments"] = json.dumps({"target_id": beta_id})
            broad_evidence = run_cu20_windows.validate_ui_evidence(
                "CU19", broad_controls, expected_workspace=workspace)
            self.assertTrue(broad_evidence["passed"], broad_evidence)

            title_only_readback = copy.deepcopy(delta)
            title_only_readback[9]["output"] = (
                f"Título de ventana: {beta.name}: Bloc de notas\n"
                f"Texto de la app (UI Automation, exacto): {beta.name}\n"
                f"1 renglón(es)\n(0.107, 0.061) {beta.name}")
            title_evidence = run_cu20_windows.validate_ui_evidence(
                "CU19", title_only_readback, expected_workspace=workspace)
            self.assertFalse(title_evidence["passed"], title_evidence)

            beta_wrong_target = list(delta)
            beta_wrong_target[8] = tool("desktop_type", {
                "target_id": "ffeedd", "text": "contenido beta solicitado"})
            rejected = run_cu20_windows.validate_ui_evidence(
                "CU19", beta_wrong_target, expected_workspace=workspace)
            self.assertFalse(rejected["passed"], rejected)

            repeated_write = [*delta[:9], delta[8], *delta[9:]]
            rejected = run_cu20_windows.validate_ui_evidence(
                "CU19", repeated_write, expected_workspace=workspace)
            self.assertFalse(rejected["passed"], rejected)

            wrong_path = list(delta)
            wrong_path[2] = tool("desktop_launch", {
                "app": "notepad", "args": f'"{alpha}"'},
                f"[desktop_launch: ok id={beta_id} opened]")
            rejected = run_cu20_windows.validate_ui_evidence(
                "CU19", wrong_path, expected_workspace=workspace)
            self.assertFalse(rejected["passed"], rejected)

            missing_post_type_read = [*delta[:9], *delta[10:]]
            rejected = run_cu20_windows.validate_ui_evidence(
                "CU19", missing_post_type_read, expected_workspace=workspace)
            self.assertFalse(rejected["passed"], rejected)

            missing_beta_title = list(delta)
            missing_beta_title[4] = tool("desktop_windows", {},
                                         f'id={alpha_id} pid=123 900x600 "{alpha.name}: Bloc de notas"')
            rejected = run_cu20_windows.validate_ui_evidence(
                "CU19", missing_beta_title, expected_workspace=workspace)
            self.assertFalse(rejected["passed"], rejected)

            missing_editor_focus = list(delta)
            missing_editor_focus[7] = tool("desktop_click_element", {
                "target_id": beta_id, "control_id": "editor-doc",
                "snapshot_id": "snap-editor-1"},
                'trace=' + json.dumps({"action": "click_element",
                                       "keyboardFocusConfirmed": False}))
            rejected = run_cu20_windows.validate_ui_evidence(
                "CU19", missing_editor_focus, expected_workspace=workspace)
            self.assertFalse(rejected["passed"], rejected)

    def test_cu19_artifact_pass_also_requires_computer_use_evidence(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            alpha, beta = run_cu20_windows.cu19_document_paths(workspace)
            alpha.write_text("", encoding="utf-8")
            beta.write_text(
                "contenido beta solicitado", encoding="utf-8")
            failed = run_cu20_windows.validate_artifact("CU19", workspace, delta=[])
            self.assertFalse(failed["passed"], failed)

    def test_calculator_requires_value_in_desktop_read(self):
        passing = [self.tool("desktop_read_text", "OCR\n42")]
        failing = [self.tool("desktop_read_text", "OCR\n24"),
                   {"role": "assistant", "content": "El resultado es 42."}]
        self.assertTrue(run_cu20_windows.validate_ui_evidence("CU01", passing)["passed"])
        self.assertFalse(run_cu20_windows.validate_ui_evidence("CU01", failing)["passed"])

    def test_browser_search_requires_visible_date_and_context(self):
        passing = [self.tool("desktop_read_text", "Apollo 11\nLanding: 20 July 1969")]
        localized = [self.tool("desktop_read_text", "Apollo 11\nFecha de aterrizaje: 20 de julio de 1969")]
        us_date = [self.tool("desktop_read_text", "Apollo 11\nLanding: July 20, 1969")]
        split_day_time_and_year = [self.tool(
            "desktop_read_text",
            "Apollo 11 landing result: landed in the Sea on July 20 at 20:17 UTC. "
            "Overview: Apollo 11 (July 16-24, 1969).")]
        failing = [self.tool("desktop_read_text", "Apollo 11 search results")]
        for read in (passing, localized, us_date, split_day_time_and_year):
            answer = {"role": "assistant", "content": "Apollo 11 landed July 20 at 20:17 UTC (1969)."}
            self.assertTrue(run_cu20_windows.validate_ui_evidence("CU05", [*read, answer])["passed"])
        missing_year = {"role": "assistant", "content": "Apollo 11 landed July 20 at 20:17 UTC."}
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU05", [*split_day_time_and_year, missing_year])["passed"])
        self.assertFalse(run_cu20_windows.validate_ui_evidence("CU05", failing)["passed"])

    def test_google_egypt_search_requires_fresh_ocr_result_click_and_open_page(self):
        target = "edge-window-1"
        url = "https://www.google.com/search?q=historia+de+Egipto+en+Wikipedia"
        typed = self.tool("desktop_type", "", json.dumps(
            {"text": url, "target_id": target}))
        bar_read = self.tool("desktop_read_text", "Google address bar query",
                             json.dumps({"target_id": target,
                                         "query": "historia+de+Egipto+en+Wikipedia"}))
        enter = self.tool("desktop_key", "", json.dumps({"key": "Enter", "target_id": target}))
        results = self.tool(
            "desktop_read_text",
            "[desktop_read_text: window 'edge-window-1']\n"
            "Título de ventana: historia de Egipto en Wikipedia - Buscar con Google\n"
            "(0.168, 0.735) Historia de Egipto - Wikipedia, la enciclopedia libre\n",
            json.dumps({"target_id": target, "query": "Wikipedia"}))
        click = self.tool("desktop_click", "[desktop_click: ok]",
                          json.dumps({"target_id": target,
                                      "x": 0.168, "y": 0.735}))
        page = self.tool("desktop_read_text",
                         "[desktop_read_text: window 'edge-window-1']\n"
                         "Título de ventana: Historia de Egipto - Wikipedia, la enciclopedia libre\n"
                         "Prehistoria de Egipto pre-3100 a. C.",
                         json.dumps({"target_id": target, "query": "Historia de Egipto"}))
        answer = {"role": "assistant",
                  "content": "La sección de Wikipedia sobre la prehistoria de Egipto la sitúa antes de 3100 a. C."}
        passing = [typed, bar_read, enter, results, click, page, answer]
        self.assertTrue(run_cu20_windows.validate_ui_evidence("CU06", passing)["passed"])
        reordered = [typed, bar_read, enter, results, click,
                     {**answer, "completedAt": 200},
                     {**page, "completedAt": 100}]
        self.assertTrue(run_cu20_windows.validate_ui_evidence("CU06", reordered)["passed"])
        premature = [typed, bar_read, enter, results, click,
                     {**answer, "completedAt": 200},
                     {**page, "completedAt": 300}]
        self.assertFalse(run_cu20_windows.validate_ui_evidence("CU06", premature)["passed"])
        no_click = [typed, bar_read, enter, results, page, answer]
        self.assertFalse(run_cu20_windows.validate_ui_evidence("CU06", no_click)["passed"])
        alternative_page = self.tool(
            "desktop_read_text",
            "[desktop_read_text: window 'edge-window-1']\n"
            "Título de ventana: Historia de Egipto - Wikipedia, la enciclopedia libre\n"
            "Narmer fue el primer faraón conocido de Egipto y habría unificado el Alto y el Bajo "
            "Egipto hacia 3150 a. C.",
            json.dumps({"target_id": target, "query": "Historia de Egipto"}))
        alternative_answer = {
            "role": "assistant",
            "content": "Según la lectura de Wikipedia, Narmer fue un faraón de Egipto hacia 3150 a. C.",
        }
        alternative_validation = run_cu20_windows.validate_ui_evidence(
            "CU06", [typed, bar_read, enter, results, click,
             alternative_page, alternative_answer])
        self.assertTrue(alternative_validation["passed"], alternative_validation)
        self.assertIn("3150", alternative_validation["visible_fact_term"])
        self.assertNotIn("3150 / 3150", alternative_validation["visible_fact_term"])
        generic_answer = {
            "role": "assistant",
            "content": "La página de Wikipedia relata la historia de Egipto.",
        }
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU06", [typed, bar_read, enter, results, click, alternative_page,
                     generic_answer])["passed"])
        wrong_coordinate = self.tool("desktop_click", "ok",
                                     json.dumps({"scope_kind": "window", "target_id": target,
                                                 "x": 0.7, "y": 0.7}))
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU06", [typed, bar_read, enter, results, wrong_coordinate, page, answer])["passed"])
        ambiguous_results = self.tool(
            "desktop_read_text",
            "[desktop_read_text: window 'edge-window-1']\n"
            "Título de ventana: historia de Egipto en Wikipedia - Buscar con Google\n"
            "(0.210, 0.021) historia de Egipto en Wikipedia - E X +\n"
            "(0.168, 0.735) Historia de Egipto - Wikipedia, la enciclopedia libre\n",
            json.dumps({"target_id": target, "query": "Wikipedia"}))
        tab_title_click = self.tool("desktop_click", "ok", json.dumps(
            {"target_id": target, "x": 0.210, "y": 0.021}))
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU06", [typed, bar_read, enter, ambiguous_results, tab_title_click, page, answer])["passed"])
        wrong_target = self.tool("desktop_click", "ok",
                                 json.dumps({"scope_kind": "window", "target_id": "other-window",
                                             "x": 0.168, "y": 0.735}))
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU06", [typed, bar_read, enter, results, wrong_target, page, answer])["passed"])
        search_page = self.tool("desktop_read_text",
                                "[desktop_read_text: window 'edge-window-1']\n"
                                "Título de ventana: historia de Egipto en Wikipedia - Buscar con Google\n"
                                "Resultado de búsqueda Wikipedia.",
                                json.dumps({"target_id": target, "query": "Historia de Egipto"}))
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU06", [typed, bar_read, enter, results, click, search_page, answer])["passed"])
        wrong_scope = self.tool("desktop_click", "ok",
                                json.dumps({"scope_kind": "screen", "target_id": target,
                                            "x": 0.168, "y": 0.735}))
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU06", [typed, bar_read, enter, results, wrong_scope, page, answer])["passed"])
        wrong_query = self.tool("desktop_type", "", json.dumps(
            {"text": "https://www.google.com/search?q=otra+busqueda", "target_id": target}))
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU06", [wrong_query, bar_read, enter, results, click, page, answer])["passed"])

    def test_find_in_page_requires_ctrl_f_and_visible_match(self):
        target = "edge-window"
        ctrl_f = self.tool("desktop_key", "ok", json.dumps({
            "key": "F", "modifiers": ["CTRL"], "target_id": target}))
        panel = self.tool("desktop_read_text", "Título de ventana: Apollo 11 - Wikipedia\n"
                          "Buscar en la página\nOCR — coordenadas (x,y):\n0 renglón(es)",
                          json.dumps({"target_id": target}))
        typed = self.tool("desktop_type", "ok", json.dumps({
            "target_id": target, "text": "July 20, 1969"}))
        typed_read = self.tool("desktop_read_text", "Buscar en la página: July 20, 1969\n"
                               "OCR — coordenadas (x,y):\n0 renglón(es)",
                               json.dumps({"target_id": target, "query": "July 20, 1969"}))
        enter = self.tool("desktop_key", "ok", json.dumps({
            "key": "ENTER", "target_id": target}))
        result = self.tool("desktop_read_text", "Buscar en la página: July 20, 1969\n"
                           "OCR — coordenadas (x,y):\n(0.42, 0.58) July 20, 1969 at 20:17 UTC",
                           json.dumps({"target_id": target, "query": "July 20, 1969"}))
        answer = {"role": "assistant", "content":
                  "Wikipedia indica que Apollo 11 alunizó el July 20, 1969 at 20:17 UTC."}
        passed = run_cu20_windows.validate_ui_evidence(
            "CU07", [ctrl_f, panel, typed, typed_read, enter, result, answer])
        self.assertTrue(passed["passed"], passed)
        spanish_answer = {"role": "assistant", "content":
                          "Apollo 11 alunizó el 20 de julio de 1969."}
        self.assertTrue(run_cu20_windows.validate_ui_evidence(
            "CU07", [ctrl_f, panel, typed, typed_read, enter, result,
                     spanish_answer])["passed"])

        immediate_match = self.tool(
            "desktop_read_text",
            "Buscar en la página: July 20, 1969\n"
            "OCR — coordenadas (x,y):\n(0.42, 0.58) July 20, 1969 at 20:17 UTC",
            json.dumps({"target_id": target, "query": "July 20, 1969"}))
        immediate = run_cu20_windows.validate_ui_evidence(
            "CU07", [ctrl_f, panel, typed, immediate_match, answer])
        self.assertTrue(immediate["passed"], immediate)
        self.assertTrue(immediate["immediate_visible_match_accepted"], immediate)

        field_only = self.tool("desktop_read_text", "Buscar en la página: July 20, 1969\n"
                               "OCR — coordenadas (x,y):\n0 renglón(es)",
                               json.dumps({"target_id": target, "query": "July 20, 1969"}))
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU07", [ctrl_f, panel, typed, typed_read, enter, field_only, answer])["passed"])
        unrelated_answer = {"role": "assistant", "content": "Apollo 11 fue una misión lunar."}
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU07", [ctrl_f, panel, typed, typed_read, enter, result,
                     unrelated_answer])["passed"])
        self.assertFalse(run_cu20_windows.validate_ui_evidence("CU07", [typed_read])["passed"])

    def test_browser_tabs_requires_two_searches_and_switch_back(self):
        messages = [
            self.tool("desktop_type", "", '{"text":"Apollo 11 landing date"}'),
            self.tool("desktop_key", "", '{"key":"Enter"}'),
            self.tool("desktop_read_text", "Apollo 11 results"),
            self.tool("desktop_key", "", '{"key":"T","modifiers":["CTRL"]}'),
            self.tool("desktop_type", "", '{"text":"Apollo 12 landing date"}'),
            self.tool("desktop_key", "", '{"key":"Enter"}'),
            self.tool("desktop_read_text", "Apollo 12 results"),
            self.tool("desktop_key", "", '{"key":"TAB","modifiers":["CTRL","SHIFT"]}'),
            self.tool("desktop_read_text", "Apollo 11 results after return"),
        ]
        validation = run_cu20_windows.validate_ui_evidence("CU08", messages)
        self.assertTrue(validation["passed"], validation)
        self.assertTrue(validation["apollo_12_visible_in_second_tab"])
        before_return = messages[:7]
        partial = run_cu20_windows.validate_ui_evidence("CU08", before_return)
        self.assertTrue(partial["apollo_12_visible_in_second_tab"])
        self.assertFalse(partial["passed"])
        without_new_tab = [m for m in messages if not (
            m.get("name") == "desktop_key" and '"key":"T"' in m.get("arguments", ""))]
        self.assertFalse(run_cu20_windows.validate_ui_evidence("CU08", without_new_tab)["passed"])
        self.assertFalse(run_cu20_windows.validate_ui_evidence("CU08", messages[:-1])["passed"])

    def cu17_sequence(self, *, selection_action="select", selected_name="CU20-select-this-file.txt",
                      click=False, repeated=False, wrong_control=False, final_read=True):
        workspace = pathlib.Path(r"C:\cu20\explorer17work")
        filename = "CU20-select-this-file.txt"
        target_id = "560b58"
        control_id = "42.31.7"
        snapshot_id = "fresh-snapshot"
        control_line = (f'controlId={control_id}  [listitem] invoke focusable=1 '
                        f'keyboard_focus=0  320x22@(400,300)  "{selected_name}"')
        launch = self.tool("desktop_launch",
                           f'[desktop_launch: ok — la ventana se abrió: id={target_id} "{workspace}"]',
                           json.dumps({"app": "explorer", "args": f'"{workspace}"'}))
        initial_read = self.tool("desktop_read_text", f"Título de ventana: {workspace}")
        query = self.tool("desktop_controls",
                          f"[desktop_controls: 1 control(es)]\nsnapshot_id={snapshot_id}\n{control_line}",
                          json.dumps({"target_id": target_id, "query": filename, "max": 8}))
        selected_control_id = "wrong-control" if wrong_control else control_id
        trace_name = selected_name
        trace = {"action": selection_action,
                 "target": {"controlId": selected_control_id, "name": trace_name,
                            "role": "listitem"}}
        select = self.tool("desktop_control_action", "[desktop_control_action: ok]\ntrace="
                           + json.dumps(trace), json.dumps({"target_id": target_id,
                           "control_id": selected_control_id, "snapshot_id": snapshot_id,
                           "action": selection_action}))
        final = self.tool("desktop_read_text",
                          f"Título de ventana: {workspace}\nNombre: {filename}",
                          json.dumps({"target_id": target_id}))
        messages = [launch, initial_read, query, select]
        if click:
            messages.append(self.tool("desktop_click", "[desktop_click: ok]",
                                      json.dumps({"target_id": target_id, "x": 0.5, "y": 0.5})))
        if repeated:
            messages.extend([query, select])
        if final_read:
            messages.append(final)
        return workspace, messages

    def test_explorer_requires_exact_ui_control_selection_and_readback(self):
        workspace, messages = self.cu17_sequence()
        result = run_cu20_windows.validate_ui_evidence("CU17", messages,
                                                       expected_workspace=workspace)
        self.assertTrue(result["passed"], result)

    def test_explorer_filename_and_random_click_is_not_a_selection(self):
        workspace, messages = self.cu17_sequence(click=True)
        result = run_cu20_windows.validate_ui_evidence("CU17", messages,
                                                       expected_workspace=workspace)
        self.assertFalse(result["passed"], result)

    def test_explorer_rejects_repeated_selection_actions(self):
        workspace, messages = self.cu17_sequence(repeated=True)
        result = run_cu20_windows.validate_ui_evidence("CU17", messages,
                                                       expected_workspace=workspace)
        self.assertFalse(result["passed"], result)

    def test_explorer_rejects_wrong_control_and_missing_readback(self):
        workspace, wrong_control = self.cu17_sequence(wrong_control=True)
        result = run_cu20_windows.validate_ui_evidence("CU17", wrong_control,
                                                       expected_workspace=workspace)
        self.assertFalse(result["passed"], result)
        workspace, no_readback = self.cu17_sequence(final_read=False)
        result = run_cu20_windows.validate_ui_evidence("CU17", no_readback,
                                                       expected_workspace=workspace)
        self.assertFalse(result["passed"], result)

    def test_explorer_prompt_requires_one_semantic_selection_without_guessing(self):
        with tempfile.TemporaryDirectory(prefix="cu17-prompt-") as temp:
            prompt = run_cu20_windows.case_prompt("CU17", pathlib.Path(temp))
            self.assertIn("desktop_controls con query='CU20-select-this-file.txt'", prompt)
            self.assertIn("action='select'", prompt)
            self.assertIn("no hagas clics, no adivines coordenadas", prompt)

    def test_cu14_requires_one_exact_write_save_and_visible_readback(self):
        line = run_cu20_windows.CU14_LINE
        target_id = "7a"
        messages = [
            self.tool("desktop_launch", '[desktop_launch: ok — la ventana se abrió: id=7a "notepad-exact-line.txt"]',
                      json.dumps({"app": "notepad", "args": r"C:\cu20\notepad-exact-line.txt"})),
            self.tool("desktop_read_text", "notepad-exact-line.txt",
                      json.dumps({"query": "notepad-exact-line.txt", "max": 8,
                                  "target_id": target_id})),
            self.tool("desktop_type", "[desktop_type: ok]", json.dumps({
                "text": line, "target_id": target_id})),
            self.tool("desktop_key", "[desktop_key: ok]", json.dumps({
                "key": "S", "modifiers": ["CTRL"], "target_id": target_id})),
            self.tool("desktop_read_text", f"1 renglón\n{line}", json.dumps({
                "query": line, "max": 8, "target_id": target_id})),
        ]
        evidence = run_cu20_windows.validate_ui_evidence("CU14", messages)
        self.assertTrue(evidence["passed"], evidence)
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            (workspace / "notepad-exact-line.txt").write_text(line, encoding="utf-8")
            result = run_cu20_windows.validate_artifact("CU14", workspace, messages)
            self.assertTrue(result["passed"], result)

        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU14", messages[:-1])["passed"])
        title_retry = list(messages)
        title_retry.insert(2, messages[1])
        self.assertTrue(run_cu20_windows.validate_ui_evidence(
            "CU14", title_retry)["passed"])

        duplicated_write = messages + [messages[2]]
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU14", duplicated_write)["passed"])
        missing_filename_confirmation = list(messages)
        del missing_filename_confirmation[1]
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU14", missing_filename_confirmation)["passed"])
        failed_filename_confirmation = [dict(message) for message in messages]
        failed_filename_confirmation[1]["ok"] = False
        self.assertFalse(run_cu20_windows.validate_ui_evidence(
            "CU14", failed_filename_confirmation)["passed"])

    def test_cu18_requires_observed_copy_paste_and_exact_saved_cell(self):
        from openpyxl import Workbook

        marker = run_cu20_windows.CLIPBOARD_MARKER
        with tempfile.TemporaryDirectory() as temp_dir:
            workspace = pathlib.Path(temp_dir)
            source = workspace / "clipboard-source.txt"
            workbook_path = workspace / "clipboard-paste.xlsx"
            source.write_text(marker, encoding="utf-8")
            workbook = Workbook()
            workbook.active["A1"] = marker
            workbook.save(workbook_path)
            notepad_id, excel_id = "ab123", "ef456a"
            box_id, box_snapshot = "name-box-42", "snapshot-name-box"
            box_controls = ("[desktop_controls: 1 control(es)]\nsnapshot_id=" + box_snapshot
                            + f'\ncontrolId={box_id}  [edit] value="A1" "Cuadro de nombres"')
            save_id, save_snapshot = "save-button-7", "snapshot-save-button"
            save_controls = ("[desktop_controls: 1 control(es)]\nsnapshot_id=" + save_snapshot
                             + f'\ncontrolId={save_id}  [button] invoke focusable=1 keyboard_focus=0 "Guardar"')
            save_trace = {"action": "click_element", "target": {
                "controlId": save_id, "name": "Guardar", "role": "button",
                "windowTargetId": excel_id}}

            def sequence(include_source_readback=False):
                events = [
                    self.tool("desktop_launch",
                              f'[desktop_launch: ok — ventana abierta: id={notepad_id} "Bloc de notas"]',
                              json.dumps({"app": "notepad", "args": f'"{source}"'})),
                    self.tool("desktop_read_text",
                              f"Título de ventana: clipboard-source.txt: Bloc de notas\nTexto visible: {marker}",
                              json.dumps({"target_id": notepad_id, "query": marker})),
                    self.tool("desktop_focus", "[desktop_focus: ok]",
                              json.dumps({"target_id": notepad_id})),
                    self.tool("desktop_key", "[desktop_key: ok]", json.dumps({
                        "key": "A", "modifiers": ["CTRL"], "target_id": notepad_id})),
                    self.tool("desktop_key", "[desktop_key: ok]", json.dumps({
                        "key": "C", "modifiers": ["CTRL"], "target_id": notepad_id})),
                    self.tool("desktop_clipboard", f"[desktop_clipboard: 27 caracteres]\n{marker}",
                              '{"op":"get"}'),
                ]
                if include_source_readback:
                    events.append(self.tool("desktop_read_text",
                        f"Título de ventana: clipboard-source.txt\n{marker}",
                        json.dumps({"target_id": notepad_id, "query": marker})))
                events.extend([
                    self.tool("desktop_clipboard", f"[desktop_clipboard: 27 caracteres]\n{marker}",
                              '{"op":"get"}'),
                    self.tool("desktop_launch",
                              f'[desktop_launch: ok — ventana abierta: id={excel_id} "Excel"]',
                              json.dumps({"app": str(run_cu20_windows.EXCEL_EXE),
                                          "args": f'/x "{workbook_path}"'})),
                    self.tool("desktop_wait_for", '[desktop_wait_for: ok]\n' + json.dumps({
                        "found": True, "elapsedMs": 1200, "windowId": excel_id,
                        "title": "clipboard-paste.xlsx - Excel"}),
                        json.dumps({"window_title": "clipboard-paste.xlsx - Excel",
                                    "timeout_ms": 30000})),
                    self.tool("desktop_read_text", "Título de ventana: clipboard-paste.xlsx - Excel",
                              json.dumps({"target_id": excel_id,
                                          "query": "clipboard-paste.xlsx"})),
                    self.tool("desktop_focus", "[desktop_focus: ok]",
                              json.dumps({"target_id": excel_id})),
                    self.tool("desktop_controls", box_controls,
                              json.dumps({"target_id": excel_id, "query": "Cuadro de nombres"})),
                    self.tool("desktop_control_action", "[desktop_control_action: ok]", json.dumps({
                        "target_id": excel_id, "control_id": box_id, "snapshot_id": box_snapshot,
                        "action": "set_value", "value": "A1"})),
                    self.tool("desktop_key", "[desktop_key: ok]", json.dumps({
                        "key": "ENTER", "target_id": excel_id})),
                    self.tool("desktop_key", "[desktop_key: ok]", json.dumps({
                        "key": "V", "modifiers": ["CTRL"], "target_id": excel_id})),
                    self.tool("desktop_controls", save_controls,
                              json.dumps({"target_id": excel_id, "query": "Guardar"})),
                    self.tool("desktop_click_element", "[desktop_click_element: ok]\ntrace="
                              + json.dumps(save_trace), json.dumps({
                                  "target_id": excel_id, "control_id": save_id,
                                  "snapshot_id": save_snapshot})),
                    self.tool("desktop_read_text",
                              f"Título de ventana: clipboard-paste.xlsx - Excel\nA1: {marker}\n"
                              "clipboard-paste.xlsx • Guardado en Este PC",
                              json.dumps({"target_id": excel_id})),
                ])
                return events

            messages = sequence()
            copy_stage = run_cu20_windows.validate_cu18_copy_stage(messages[:6], workspace)
            self.assertTrue(copy_stage["passed"], copy_stage)
            result = run_cu20_windows.validate_artifact("CU18", workspace, messages)
            self.assertTrue(result["passed"], result)
            self.assertEqual(result["A1"], marker)
            self.assertEqual(result["nonempty_cells"], 1)

            optional_readback = sequence(include_source_readback=True)
            optional_result = run_cu20_windows.validate_artifact(
                "CU18", workspace, optional_readback)
            self.assertTrue(optional_result["passed"], optional_result)

            disabled_save = list(messages)
            disabled_save[15] = self.tool(
                "desktop_controls", save_controls.replace("focusable=1", "enabled=0 focusable=0"),
                json.dumps({"target_id": excel_id, "query": "Guardar"}))
            self.assertFalse(run_cu20_windows.validate_ui_evidence(
                "CU18", disabled_save, expected_workspace=workspace)["passed"])

            wrong_second_clipboard = list(messages)
            wrong_second_clipboard[6] = self.tool(
                "desktop_clipboard", "[desktop_clipboard: 2 caracteres]\nA1", '{"op":"get"}')
            self.assertFalse(run_cu20_windows.validate_artifact(
                "CU18", workspace, wrong_second_clipboard)["passed"])

            wrong_control = list(messages)
            action_args = json.loads(wrong_control[12]["arguments"])
            action_args["control_id"] = "unrelated-control"
            wrong_control[12]["arguments"] = json.dumps(action_args)
            self.assertFalse(run_cu20_windows.validate_ui_evidence(
                "CU18", wrong_control, expected_workspace=workspace)["passed"])

            missing_saved_readback = list(messages[:-1])
            self.assertFalse(run_cu20_windows.validate_artifact(
                "CU18", workspace, missing_saved_readback)["passed"])

            duplicated_paste = list(messages)
            duplicated_paste.insert(15, self.tool(
                "desktop_key", "[desktop_key: ok]",
                '{"key":"V","modifiers":["CTRL"],"target_id":"ef456a"}'))
            self.assertFalse(run_cu20_windows.validate_ui_evidence(
                "CU18", duplicated_paste, expected_workspace=workspace)["passed"])

            (workspace / "clipboard-source.txt").write_text("tampered", encoding="utf-8")
            tampered_source = run_cu20_windows.validate_artifact("CU18", workspace, messages)
            self.assertFalse(tampered_source["passed"], tampered_source)

    def test_firewall_prompt_stops_agent_before_followup_action(self):
        delta = [{"role": "toolcall", "name": "desktop_read_text",
                  "output": "Firewall de Windows bloqueó algunas características de test_system_profiles.exe."}]
        with patch.object(run_cu20_windows, "invoke") as invoke:
            with self.assertRaisesRegex(run_cu20_windows.TurnSafetyError, "diálogo de permiso"):
                run_cu20_windows.enforce_no_windows_permission_prompt(delta)
        invoke.assert_called_once_with("stopAgent")

    def test_normal_inprivate_read_does_not_stop_agent(self):
        delta = [{"role": "toolcall", "name": "desktop_read_text",
                  "output": "Nueva pestaña InPrivate. Dirección y barra de búsqueda."}]
        with patch.object(run_cu20_windows, "invoke") as invoke:
            run_cu20_windows.enforce_no_windows_permission_prompt(delta)
        invoke.assert_not_called()

    def test_firewall_prompt_in_launch_receipt_stops_agent(self):
        delta = [{"role": "toolcall", "name": "desktop_launch",
                  "output": "Windows Defender Firewall blocked some features of app.exe."}]
        with patch.object(run_cu20_windows, "invoke") as invoke:
            with self.assertRaisesRegex(run_cu20_windows.TurnSafetyError, "diálogo de permiso"):
                run_cu20_windows.enforce_no_windows_permission_prompt(delta)
        invoke.assert_called_once_with("stopAgent")


if __name__ == "__main__":
    unittest.main()
